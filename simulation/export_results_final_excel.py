import os
import re
import numpy as np
import pandas as pd
import dispaset as ds

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# 1. PATHS
# ============================================================

SIMULATION_PATH = r'../Simulations/simulation_(2)_BE-NUC'

BOUNDARY_DATA_PATH = (
    r'../Database/Boundary_Sector/'
    r'BoundarySectorData/BE/2023.CSV'
)

BOUNDARY_DEMAND_PATH = (
    r'../Database/Boundary_Sector/'
    r'BoundarySectorDemand/BE/2023.CSV'
)

# IMPORTANT:
# This is the SAME final Excel file.
# Running the script again will overwrite it.
OUTPUT_FILE = os.path.join(
    SIMULATION_PATH,
    'UNLEASH_BE_NUC_FINAL_RESULTS.xlsx'
)


# ============================================================
# 2. LOAD DISPA-SET RESULTS
# ============================================================

print()
print("=" * 65)
print("LOADING UNLEASH (2)_BE-NUC SIMULATION RESULTS")
print("=" * 65)

inputs, results = ds.get_sim_results(
    path=SIMULATION_PATH
)

print("Results loaded successfully.")
print("Number of Dispa-SET result variables:", len(results))


# ============================================================
# 3. COMMON HOURLY INDEX
# ============================================================

HOURLY_INDEX = pd.date_range(
    start='2023-01-01 00:00:00',
    periods=8760,
    freq='h'
)

HOURLY_INDEX.name = 'Date / Time'


# ============================================================
# 4. HELPER FUNCTIONS
# ============================================================

def clean_small_values(obj, tolerance=1e-9):
    """
    Replace very small numerical noise with zero.
    """

    if isinstance(obj, pd.Series):

        obj = obj.copy()

        numeric = pd.to_numeric(
            obj,
            errors='coerce'
        )

        mask = numeric.abs() < tolerance

        obj.loc[mask.fillna(False)] = 0

        return obj


    if isinstance(obj, pd.DataFrame):

        obj = obj.copy()

        numeric_cols = obj.select_dtypes(
            include=[np.number]
        ).columns

        for col in numeric_cols:

            mask = (
                obj[col].abs()
                < tolerance
            )

            obj.loc[mask, col] = 0

        return obj


    return obj


def get_result_df(name):
    """
    Return one Dispa-SET result as a cleaned DataFrame.

    All full-year hourly results are forced onto the same
    2023 hourly Date / Time index.
    """

    if name not in results:
        return None

    data = results[name]

    if isinstance(data, pd.Series):
        data = data.to_frame()

    if not isinstance(data, pd.DataFrame):
        return None

    if data.shape[1] == 0:
        return None

    df = data.copy()


    if len(df) == 8760:

        df.index = HOURLY_INDEX.copy()
        df.index.name = 'Date / Time'


    elif isinstance(
        df.index,
        pd.DatetimeIndex
    ):

        if df.index.tz is not None:

            df.index = (
                df.index
                .tz_localize(None)
            )

        df.index.name = 'Date / Time'


    return clean_small_values(df)


def first_result_column(
    name,
    new_name
):
    """
    Extract the first column from a Dispa-SET result.
    """

    df = get_result_df(name)

    if df is None or df.empty:
        return None

    series = df.iloc[:, 0].copy()

    series.name = new_name
    series.index.name = 'Date / Time'

    return clean_small_values(
        series
    )


def add_series(
    table,
    series
):
    """
    Join a Series to an hourly table without losing
    the Date / Time index name.
    """

    if series is None:
        return table


    if table is None:

        table = pd.DataFrame(
            index=series.index
        )

        table.index.name = (
            'Date / Time'
        )


    result = table.join(
        series,
        how='outer'
    )

    result.index.name = (
        'Date / Time'
    )

    return result


def annual_sum(name):
    """
    Sum a Dispa-SET result over the entire year,
    separately for each unit/column.
    """

    df = get_result_df(name)

    if df is None:
        return pd.Series(dtype=float)

    return (
        df.fillna(0)
        .sum(axis=0)
    )


def scalar_result(name):
    """
    Extract a scalar result.
    """

    df = get_result_df(name)

    if df is None or df.empty:
        return 0.0

    try:

        return float(
            df.iloc[0, 0]
        )

    except Exception:

        return 0.0


def safe_sum(series):
    """
    Sum a Series safely.
    """

    if series is None:
        return 0.0

    return float(
        pd.to_numeric(
            series,
            errors='coerce'
        )
        .fillna(0)
        .sum()
    )


def read_single_series_csv(
    path,
    series_name
):
    """
    Read a two-column hourly CSV
    and return its numeric column.
    """

    try:

        df = pd.read_csv(path)

        if len(df) != 8760:

            print(
                f"Warning: {path} does not "
                "contain 8760 rows."
            )

            return None


        numeric_cols = []

        for col in df.columns:

            converted = pd.to_numeric(
                df[col],
                errors='coerce'
            )

            if (
                converted
                .notna()
                .sum()
                > 8000
            ):

                numeric_cols.append(col)


        if not numeric_cols:
            return None


        series = pd.to_numeric(
            df[numeric_cols[-1]],
            errors='coerce'
        )


        series.index = (
            HOURLY_INDEX.copy()
        )

        series.index.name = (
            'Date / Time'
        )

        series.name = (
            series_name
        )

        return series


    except Exception as error:

        print(
            f"Warning: could not read "
            f"{path}: {error}"
        )

        return None


def get_unit_column(
    unit_df,
    possible_names
):
    """
    Find a unit-data column using several possible names.
    """

    for name in possible_names:

        if name in unit_df.columns:
            return name

    return None


# ============================================================
# 5. DISPA-SET UNIT DATA
# ============================================================

if (
    'units' in inputs
    and isinstance(
        inputs['units'],
        pd.DataFrame
    )
):

    units = inputs[
        'units'
    ].copy()

else:

    units = pd.DataFrame()


units.index.name = 'Unit'


# ============================================================
# 6. BOUNDARY SECTOR DEMAND
# ============================================================

boundary_demand = (
    read_single_series_csv(
        BOUNDARY_DEMAND_PATH,
        'Boundary Sector Demand [MW]'
    )
)


# ============================================================
# 7. BOUNDARY SECTOR PARAMETERS
# ============================================================

storage_capacity = 0.0
storage_self_discharge_rate = 0.0
minimum_soc = 0.0
vpp_max_power = 0.0
cost_x_not_served = 0.0


try:

    boundary_data = pd.read_csv(
        BOUNDARY_DATA_PATH
    )


    if (
        'STOCapacity'
        in boundary_data.columns
    ):

        storage_capacity = float(
            boundary_data[
                'STOCapacity'
            ].iloc[0]
        )


    if (
        'STOSelfDischarge'
        in boundary_data.columns
    ):

        storage_self_discharge_rate = float(
            boundary_data[
                'STOSelfDischarge'
            ].iloc[0]
        )


    if (
        'STOMinSOC'
        in boundary_data.columns
    ):

        minimum_soc = float(
            boundary_data[
                'STOMinSOC'
            ].iloc[0]
        )


    if (
        'STOMaxPower'
        in boundary_data.columns
    ):

        vpp_max_power = float(
            boundary_data[
                'STOMaxPower'
            ].iloc[0]
        )


    if (
        'CostXNotServed'
        in boundary_data.columns
    ):

        cost_x_not_served = float(
            boundary_data[
                'CostXNotServed'
            ].iloc[0]
        )


except Exception as error:

    print(
        "Warning: Boundary Sector "
        "parameters could not be loaded:"
    )

    print(error)


# ============================================================
# 8. VPP CONVERSION EFFICIENCY
# ============================================================

vpp_efficiency = 0.0


if not units.empty:

    efficiency_sector_col = (
        get_unit_column(
            units,
            [
                'EfficiencySector1',
                'EfficiencySector'
            ]
        )
    )


    if efficiency_sector_col:

        vpp_rows = [
            idx
            for idx
            in units.index
            if 'VPP'
            in str(idx).upper()
        ]


        if vpp_rows:

            value = units.loc[
                vpp_rows[0],
                efficiency_sector_col
            ]


            if pd.notna(value):

                vpp_efficiency = float(
                    value
                )


# ============================================================
# 9. BASIC HOURLY RESULTS
# ============================================================

total_demand = first_result_column(
    'TotalDemand',
    'Total Electricity Demand incl. VPP [MW]'
)


vpp_consumption = first_result_column(
    'OutputPowerConsumption',
    'VPP Electricity Consumption [MW]'
)


vpp_power_x = first_result_column(
    'OutputPowerX',
    'VPP Useful PowerX [MW]'
)


# Base electricity demand is derived from the model balance:
# Total Demand = Base Demand + VPP Electricity Consumption

base_load = None


if (
    total_demand is not None
    and vpp_consumption is not None
):

    base_load = (
        total_demand
        - vpp_consumption
    )

    base_load.name = (
        'Base Electricity Demand [MW]'
    )

    base_load.index.name = (
        'Date / Time'
    )


# ============================================================
# 10. DISPATCH TABLE
# ============================================================

dispatch = pd.DataFrame(
    index=HOURLY_INDEX.copy()
)

dispatch.index.name = (
    'Date / Time'
)


dispatch = add_series(
    dispatch,
    base_load
)


dispatch = add_series(
    dispatch,
    total_demand
)


residual_load = first_result_column(
    'OutputResidualLoad',
    'Residual Load [MW]'
)


dispatch = add_series(
    dispatch,
    residual_load
)


electric_storage_input = (
    first_result_column(
        'OutputStorageInput',
        'Electric Storage Charging [MW]'
    )
)


dispatch = add_series(
    dispatch,
    electric_storage_input
)


electric_storage_level = (
    first_result_column(
        'OutputStorageLevel',
        'Electric Storage Level [MWh]'
    )
)


dispatch = add_series(
    dispatch,
    electric_storage_level
)


power = get_result_df(
    'OutputPower'
)


if power is not None:

    power = power.copy()

    power.columns = [
        f'Generation | {unit} [MW]'
        for unit
        in power.columns
    ]


    dispatch = dispatch.join(
        power,
        how='left'
    )

    dispatch.index.name = (
        'Date / Time'
    )


curtailment = get_result_df(
    'OutputCurtailmentPerUnit'
)


if curtailment is not None:

    total_curtailment = (
        curtailment
        .fillna(0)
        .sum(axis=1)
    )

    total_curtailment.name = (
        'Total Curtailment [MW]'
    )

    total_curtailment.index.name = (
        'Date / Time'
    )

else:

    total_curtailment = pd.Series(
        0.0,
        index=HOURLY_INDEX.copy(),
        name='Total Curtailment [MW]'
    )

    total_curtailment.index.name = (
        'Date / Time'
    )


dispatch = add_series(
    dispatch,
    total_curtailment
)


dispatch = clean_small_values(
    dispatch
)


# ============================================================
# 11. VPP / BOUNDARY SECTOR TABLE
# ============================================================

vpp = pd.DataFrame(
    index=HOURLY_INDEX.copy()
)

vpp.index.name = (
    'Date / Time'
)


vpp = add_series(
    vpp,
    vpp_consumption
)


vpp = add_series(
    vpp,
    vpp_power_x
)


vpp = add_series(
    vpp,
    boundary_demand
)


sector_storage_input = (
    first_result_column(
        'OutputSectorXStorageInput',
        'SectorX Storage Net Input [MW]'
    )
)


vpp = add_series(
    vpp,
    sector_storage_input
)


sector_soc = first_result_column(
    'OutputSectorXStorageLevel',
    'SectorX Storage SOC [%]'
)


vpp = add_series(
    vpp,
    sector_soc
)


# Stored energy derived from:
# SOC x Storage Capacity

if (
    sector_soc is not None
    and storage_capacity > 0
):

    stored_energy = (
        sector_soc
        * storage_capacity
    )

    stored_energy.name = (
        'SectorX Stored Energy [MWh]'
    )

    stored_energy.index.name = (
        'Date / Time'
    )


    vpp = add_series(
        vpp,
        stored_energy
    )


sector_self_discharge = (
    first_result_column(
        'OutputSectorXSelfDischarge',
        'SectorX Self-Discharge [MW]'
    )
)


vpp = add_series(
    vpp,
    sector_self_discharge
)


sector_storage_shadow = (
    first_result_column(
        'OutputSectorXStorageShadowPrice',
        'SectorX Storage Shadow Price [EUR/MWh]'
    )
)


vpp = add_series(
    vpp,
    sector_storage_shadow
)


sector_shadow = (
    first_result_column(
        'SectorXShadowPrice',
        'SectorX Shadow Price [EUR/MWh]'
    )
)


vpp = add_series(
    vpp,
    sector_shadow
)


vpp = clean_small_values(
    vpp
)


# ============================================================
# 12. SYSTEM TABLE
# ============================================================

system = pd.DataFrame(
    index=HOURLY_INDEX.copy()
)

system.index.name = (
    'Date / Time'
)


system = add_series(
    system,
    base_load
)


system = add_series(
    system,
    total_demand
)


system = add_series(
    system,
    residual_load
)


system_cost_hourly = (
    first_result_column(
        'OutputSystemCost',
        'System Cost [EUR]'
    )
)


system = add_series(
    system,
    system_cost_hourly
)


emissions_hourly = (
    first_result_column(
        'OutputEmissions',
        'CO2 Emissions [tCO2]'
    )
)


system = add_series(
    system,
    emissions_hourly
)


electricity_price = (
    first_result_column(
        'ShadowPrice',
        'Electricity Shadow Price [EUR/MWh]'
    )
)


system = add_series(
    system,
    electricity_price
)


load_shedding = first_result_column(
    'OutputShedLoad',
    'Load Shedding [MW]'
)


system = add_series(
    system,
    load_shedding
)


energy_not_served = (
    first_result_column(
        'ENSH-EnergyNotServedHourly',
        'Energy Not Served [MWh]'
    )
)


system = add_series(
    system,
    energy_not_served
)


system = add_series(
    system,
    total_curtailment
)


system = clean_small_values(
    system
)


# ============================================================
# 13. RESERVES TABLE
# ============================================================

reserves = pd.DataFrame(
    index=HOURLY_INDEX.copy()
)

reserves.index.name = (
    'Date / Time'
)


reserve_variables = [

    (
        'OutputDemand_2U',
        '2U Reserve Requirement [MW]'
    ),

    (
        'OutputDemand_2D',
        '2D Reserve Requirement [MW]'
    ),

    (
        'OutputDemand_3U',
        '3U Reserve Requirement [MW]'
    ),

    (
        'ShadowPrice_2U',
        '2U Reserve Shadow Price [EUR/MW]'
    ),

    (
        'ShadowPrice_2D',
        '2D Reserve Shadow Price [EUR/MW]'
    ),

    (
        'ShadowPrice_3U',
        '3U Reserve Shadow Price [EUR/MW]'
    ),
]


for (
    original_name,
    clean_name
) in reserve_variables:

    reserves = add_series(

        reserves,

        first_result_column(
            original_name,
            clean_name
        )
    )


for (
    source_name,
    clean_name
) in [

    (
        'OutputReserve_2U',
        'Available 2U Reserve [MW]'
    ),

    (
        'OutputReserve_2D',
        'Available 2D Reserve [MW]'
    ),

    (
        'OutputReserve_3U',
        'Available 3U Reserve [MW]'
    ),
]:

    reserve_df = get_result_df(
        source_name
    )


    if reserve_df is not None:

        total = (
            reserve_df
            .fillna(0)
            .sum(axis=1)
        )

        total.name = clean_name

        total.index.name = (
            'Date / Time'
        )


        reserves = add_series(
            reserves,
            total
        )


reserves = clean_small_values(
    reserves
)


# ============================================================
# 14. UNITS SUMMARY
# ============================================================

generation = annual_sum(
    'OutputPower'
)


committed = annual_sum(
    'OutputCommitted'
)


variable_cost = annual_sum(
    'UnitHourlyVariableCost'
)


startup_cost = annual_sum(
    'UnitHourlyStartUpCost'
)


# NOTE:
# UnitHourlyProductionCost is intentionally NOT used.
#
# It produced the same non-zero annual value for units
# with zero generation and zero commitment, including VPP_BE.
# Therefore it is not presented as a physical unit-specific
# production cost in the final workbook.


unit_summary = pd.DataFrame(
    index=units.index
)

unit_summary.index.name = (
    'Unit'
)


# ------------------------------------------------------------
# UNIT CHARACTERISTICS
# ------------------------------------------------------------

characteristics = {

    'Zone':
        ['Zone'],

    'Technology':
        ['Technology'],

    'Fuel':
        ['Fuel'],

    'Installed Capacity [MW]':
        [
            'PowerCapacity',
            'Power Capacity',
            'Capacity'
        ],

    'Efficiency [-]':
        [
            'Efficiency'
        ],

    'Emission Rate [tCO2/MWh]':
        [
            'EmissionRate',
            'CO2Intensity'
        ],

    'Storage Capacity [MWh]':
        [
            'StorageCapacity',
            'STOCapacity'
        ],
}


for (
    new_name,
    possible_names
) in characteristics.items():

    col = get_unit_column(
        units,
        possible_names
    )


    if col:

        unit_summary[
            new_name
        ] = units[col]


# ------------------------------------------------------------
# ANNUAL UNIT RESULTS
# ------------------------------------------------------------

unit_summary[
    'Annual Generation [MWh]'
] = generation.reindex(
    unit_summary.index
).fillna(0)


unit_summary[
    'Committed Hours [h]'
] = committed.reindex(
    unit_summary.index
).fillna(0)


unit_summary[
    'Annual Variable Cost [EUR]'
] = variable_cost.reindex(
    unit_summary.index
).fillna(0)


unit_summary[
    'Annual Startup Cost [EUR]'
] = startup_cost.reindex(
    unit_summary.index
).fillna(0)


# ------------------------------------------------------------
# CAPACITY FACTOR
# ------------------------------------------------------------

if (
    'Installed Capacity [MW]'
    in unit_summary.columns
):

    installed_capacity = (
        pd.to_numeric(
            unit_summary[
                'Installed Capacity [MW]'
            ],
            errors='coerce'
        )
        .fillna(0)
    )


    denominator = (
        installed_capacity
        * 8760
    )


    unit_summary[
        'Capacity Factor [%]'
    ] = np.where(

        denominator > 0,

        unit_summary[
            'Annual Generation [MWh]'
        ]
        / denominator,

        0
    )


# ------------------------------------------------------------
# ANNUAL CO2 EMISSIONS
# ------------------------------------------------------------

if (
    'Emission Rate [tCO2/MWh]'
    in unit_summary.columns
):

    emission_rate = (
        pd.to_numeric(
            unit_summary[
                'Emission Rate [tCO2/MWh]'
            ],
            errors='coerce'
        )
        .fillna(0)
    )


    unit_summary[
        'Annual CO2 Emissions [tCO2]'
    ] = (

        emission_rate
        *
        unit_summary[
            'Annual Generation [MWh]'
        ]
    )


unit_summary = clean_small_values(
    unit_summary
)


# ============================================================
# 15. OVERVIEW KPIs
# ============================================================

annual_base_demand = safe_sum(
    base_load
)


annual_total_demand = safe_sum(
    total_demand
)


annual_vpp_consumption = safe_sum(
    vpp_consumption
)


annual_power_x = safe_sum(
    vpp_power_x
)


annual_boundary_demand = safe_sum(
    boundary_demand
)


annual_self_discharge = safe_sum(
    sector_self_discharge
)


annual_system_cost = safe_sum(
    system_cost_hourly
)


annual_emissions = safe_sum(
    emissions_hourly
)


annual_curtailment = safe_sum(
    total_curtailment
)


annual_load_shedding = safe_sum(
    load_shedding
)


annual_ens = safe_sum(
    energy_not_served
)


lolh = scalar_result(
    'LOLH-LosOfLoadHours'
)


lolp = scalar_result(
    'LOLP-LosOfLoadProbability'
)


lolf = scalar_result(
    'LOLF-LosOfLoadFrequency'
)


overview_kpis = pd.DataFrame(

    [

        [
            'Simulated Hours',
            8760,
            'h'
        ],

        [
            'Base Electricity Demand',
            annual_base_demand,
            'MWh'
        ],

        [
            'Total Electricity Demand incl. VPP',
            annual_total_demand,
            'MWh'
        ],

        [
            'VPP Electricity Consumption',
            annual_vpp_consumption,
            'MWh'
        ],

        [
            'VPP Useful PowerX',
            annual_power_x,
            'MWh'
        ],

        [
            'Boundary Sector Demand',
            annual_boundary_demand,
            'MWh'
        ],

        [
            'SectorX Self-Discharge',
            annual_self_discharge,
            'MWh'
        ],

        [
            'Total System Cost',
            annual_system_cost,
            'EUR'
        ],

        [
            'Total CO2 Emissions',
            annual_emissions,
            'tCO2'
        ],

        [
            'Total Curtailment',
            annual_curtailment,
            'MWh'
        ],

        [
            'Total Load Shedding',
            annual_load_shedding,
            'MWh'
        ],

        [
            'Energy Not Served',
            annual_ens,
            'MWh'
        ],

        [
            'Loss of Load Hours',
            lolh,
            'h'
        ],

        [
            'Loss of Load Probability',
            lolp,
            '-'
        ],

        [
            'Loss of Load Frequency',
            lolf,
            '-'
        ],
    ],

    columns=[
        'Key Performance Indicator',
        'Value',
        'Unit'
    ]
)


# ============================================================
# 16. OVERVIEW SCENARIO PARAMETERS
# ============================================================

overview_parameters = pd.DataFrame(

    [

        [
            'Scenario',
            '(2)_BE-NUC'
        ],

        [
            'Zone',
            'Belgium (BE)'
        ],

        [
            'Simulation Year',
            '2023'
        ],

        [
            'Time Resolution',
            '1 hour'
        ],

        [
            'VPP Maximum Power',
            f'{vpp_max_power:,.0f} MW'
        ],

        [
            'VPP Conversion Efficiency',
            f'{vpp_efficiency:.1%}'
        ],

        [
            'SectorX Storage Capacity',
            f'{storage_capacity:,.0f} MWh'
        ],

        [
            'SectorX Minimum SOC',
            f'{minimum_soc:.1%}'
        ],

        [
            'SectorX Self-Discharge Rate',
            (
                f'{storage_self_discharge_rate:.4%} '
                'per hour'
            )
        ],

        [
            'Cost of Unserved SectorX Demand',
            (
                f'{cost_x_not_served:,.0f} '
                'EUR/MWh'
            )
        ],
    ],

    columns=[
        'Scenario Parameter',
        'Value'
    ]
)


# ============================================================
# 17. RESET INDEX FOR EXCEL
# ============================================================

dispatch_excel = (
    dispatch
    .reset_index()
)


vpp_excel = (
    vpp
    .reset_index()
)


system_excel = (
    system
    .reset_index()
)


reserves_excel = (
    reserves
    .reset_index()
)


unit_summary_excel = (
    unit_summary
    .reset_index()
)


# Force the first column names explicitly

for df in [
    dispatch_excel,
    vpp_excel,
    system_excel,
    reserves_excel
]:

    df.rename(
        columns={
            df.columns[0]:
            'Date / Time'
        },
        inplace=True
    )


unit_summary_excel.rename(
    columns={
        unit_summary_excel.columns[0]:
        'Unit'
    },
    inplace=True
)


# ============================================================
# 18. EXPORT DATA TO THE SAME FINAL EXCEL FILE
# ============================================================

print()
print("Creating / overwriting final Excel workbook...")


with pd.ExcelWriter(
    OUTPUT_FILE,
    engine='openpyxl'
) as writer:


    overview_kpis.to_excel(
        writer,
        sheet_name='Overview',
        index=False,
        startrow=4,
        startcol=0
    )


    overview_parameters.to_excel(
        writer,
        sheet_name='Overview',
        index=False,
        startrow=4,
        startcol=5
    )


    dispatch_excel.to_excel(
        writer,
        sheet_name='Dispatch',
        index=False,
        startrow=3
    )


    vpp_excel.to_excel(
        writer,
        sheet_name='VPP_Boundary',
        index=False,
        startrow=3
    )


    system_excel.to_excel(
        writer,
        sheet_name='System',
        index=False,
        startrow=3
    )


    reserves_excel.to_excel(
        writer,
        sheet_name='Reserves',
        index=False,
        startrow=3
    )


    unit_summary_excel.to_excel(
        writer,
        sheet_name='Units_Summary',
        index=False,
        startrow=3
    )


# ============================================================
# 19. PROFESSIONAL FORMATTING
# ============================================================

print("Applying professional formatting...")


wb = load_workbook(
    OUTPUT_FILE
)


# ------------------------------------------------------------
# COLOR PALETTE
# ------------------------------------------------------------

NAVY = '17365D'
BLUE = '1F4E78'
MEDIUM_BLUE = '5B9BD5'
VERY_LIGHT_BLUE = 'EEF5FB'

GREEN = '548235'
LIGHT_GREEN = 'E2F0D9'

ORANGE = 'ED7D31'

LIGHT_RED = 'F4CCCC'

WHITE = 'FFFFFF'


thin_border = Border(

    bottom=Side(
        style='thin',
        color='D9E1F2'
    )
)


# ============================================================
# TITLE FUNCTION
# ============================================================

def add_title(
    ws,
    title_text,
    subtitle_text,
    color=NAVY
):

    max_col = max(
        ws.max_column,
        4
    )


    end_letter = get_column_letter(
        max_col
    )


    ws.merge_cells(
        f'A1:{end_letter}1'
    )


    title = ws['A1']

    title.value = (
        title_text
    )

    title.fill = PatternFill(
        'solid',
        fgColor=color
    )

    title.font = Font(
        color=WHITE,
        bold=True,
        size=17
    )

    title.alignment = Alignment(
        horizontal='center',
        vertical='center'
    )

    ws.row_dimensions[1].height = (
        34
    )


    ws.merge_cells(
        f'A2:{end_letter}2'
    )


    subtitle = ws['A2']

    subtitle.value = (
        subtitle_text
    )

    subtitle.fill = PatternFill(
        'solid',
        fgColor=MEDIUM_BLUE
    )

    subtitle.font = Font(
        color=WHITE,
        italic=True,
        size=10
    )

    subtitle.alignment = Alignment(
        horizontal='center',
        vertical='center'
    )

    ws.row_dimensions[2].height = (
        24
    )


# ============================================================
# STANDARD DATA SHEET FORMAT
# ============================================================

def format_data_sheet(
    ws,
    title,
    subtitle,
    header_color=NAVY
):

    add_title(
        ws,
        title,
        subtitle,
        header_color
    )


    header_row = 4


    # --------------------------------------------------------
    # HEADERS
    # --------------------------------------------------------

    for cell in ws[header_row]:

        cell.fill = PatternFill(
            'solid',
            fgColor=header_color
        )

        cell.font = Font(
            color=WHITE,
            bold=True,
            size=9
        )

        cell.alignment = Alignment(
            horizontal='center',
            vertical='center',
            wrap_text=True
        )

        cell.border = (
            thin_border
        )


    ws.row_dimensions[
        header_row
    ].height = 42


    # Freeze first column and header row

    ws.freeze_panes = 'B5'


    # --------------------------------------------------------
    # BODY
    # --------------------------------------------------------

    for row in ws.iter_rows(
        min_row=5
    ):

        for cell in row:

            cell.border = (
                thin_border
            )

            cell.alignment = Alignment(
                vertical='center'
            )


            if isinstance(
                cell.value,
                (int, float)
            ):

                cell.number_format = (
                    '#,##0.00'
                )


    # --------------------------------------------------------
    # DATE/TIME FORMAT
    # --------------------------------------------------------

    if (
        ws.cell(
            row=4,
            column=1
        ).value
        == 'Date / Time'
    ):

        for row in range(
            5,
            ws.max_row + 1
        ):

            ws.cell(
                row=row,
                column=1
            ).number_format = (
                'dd/mm/yyyy hh:mm'
            )


    # --------------------------------------------------------
    # EXCEL TABLE
    # --------------------------------------------------------

    safe_name = re.sub(
        r'[^A-Za-z0-9_]',
        '_',
        ws.title
    )


    table_ref = (

        f'A4:'
        f'{get_column_letter(ws.max_column)}'
        f'{ws.max_row}'
    )


    table = Table(
        displayName=(
            f'Table_{safe_name}'
        ),
        ref=table_ref
    )


    style = TableStyleInfo(
        name='TableStyleMedium2',
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )


    table.tableStyleInfo = (
        style
    )


    ws.add_table(
        table
    )


    # --------------------------------------------------------
    # COLUMN WIDTHS
    # --------------------------------------------------------

    for col_idx in range(
        1,
        ws.max_column + 1
    ):

        header = str(
            ws.cell(
                row=4,
                column=col_idx
            ).value
        )


        letter = get_column_letter(
            col_idx
        )


        if (
            header
            == 'Date / Time'
        ):

            width = 21


        elif (
            header
            == 'Unit'
        ):

            width = 42


        elif (
            'Generation |'
            in header
        ):

            width = 34


        elif (
            'Shadow Price'
            in header
        ):

            width = 25


        elif (
            'Cost'
            in header
        ):

            width = 25


        elif (
            'Emission'
            in header
        ):

            width = 24


        elif (
            'Capacity'
            in header
        ):

            width = 23


        else:

            width = min(
                max(
                    len(header) + 3,
                    17
                ),
                30
            )


        ws.column_dimensions[
            letter
        ].width = width


# ============================================================
# 20. FORMAT OVERVIEW
# ============================================================

ws = wb[
    'Overview'
]


ws.merge_cells(
    'A1:H1'
)


ws['A1'] = (
    'UNLEASH - (2)_BE-NUC Final Simulation Results'
)


ws['A1'].fill = PatternFill(
    'solid',
    fgColor=NAVY
)


ws['A1'].font = Font(
    color=WHITE,
    bold=True,
    size=18
)


ws['A1'].alignment = Alignment(
    horizontal='center',
    vertical='center'
)


ws.row_dimensions[1].height = (
    36
)


ws.merge_cells(
    'A2:H2'
)


ws['A2'] = (
    'Belgian Boundary Sector / '
    'Virtual Power Plant - 2023'
)


ws['A2'].fill = PatternFill(
    'solid',
    fgColor=MEDIUM_BLUE
)


ws['A2'].font = Font(
    color=WHITE,
    italic=True,
    size=11
)


ws['A2'].alignment = Alignment(
    horizontal='center',
    vertical='center'
)


# ------------------------------------------------------------
# SECTION TITLES
# ------------------------------------------------------------

ws['A4'] = (
    'KEY PERFORMANCE INDICATORS'
)


ws['A4'].fill = PatternFill(
    'solid',
    fgColor=BLUE
)


ws['A4'].font = Font(
    color=WHITE,
    bold=True
)


ws['F4'] = (
    'SCENARIO PARAMETERS'
)


ws['F4'].fill = PatternFill(
    'solid',
    fgColor=GREEN
)


ws['F4'].font = Font(
    color=WHITE,
    bold=True
)


# ------------------------------------------------------------
# TABLE HEADERS
# ------------------------------------------------------------

for cell in ws[5]:

    if cell.column <= 3:

        cell.fill = PatternFill(
            'solid',
            fgColor=BLUE
        )

        cell.font = Font(
            color=WHITE,
            bold=True
        )


    if (
        6
        <= cell.column
        <= 7
    ):

        cell.fill = PatternFill(
            'solid',
            fgColor=GREEN
        )

        cell.font = Font(
            color=WHITE,
            bold=True
        )


# ------------------------------------------------------------
# KPI ROWS
# ------------------------------------------------------------

for row in range(
    6,
    ws.max_row + 1
):

    for col in range(
        1,
        4
    ):

        cell = ws.cell(
            row,
            col
        )

        cell.border = (
            thin_border
        )


        if row % 2 == 0:

            cell.fill = PatternFill(
                'solid',
                fgColor=VERY_LIGHT_BLUE
            )


    for col in range(
        6,
        8
    ):

        cell = ws.cell(
            row,
            col
        )

        cell.border = (
            thin_border
        )


        if row % 2 == 0:

            cell.fill = PatternFill(
                'solid',
                fgColor=LIGHT_GREEN
            )


# KPI number formatting

for row in range(
    6,
    6 + len(
        overview_kpis
    )
):

    ws.cell(
        row,
        2
    ).number_format = (
        '#,##0.00'
    )


ws.column_dimensions['A'].width = 38
ws.column_dimensions['B'].width = 22
ws.column_dimensions['C'].width = 15

ws.column_dimensions['D'].width = 4
ws.column_dimensions['E'].width = 4

ws.column_dimensions['F'].width = 38
ws.column_dimensions['G'].width = 28
ws.column_dimensions['H'].width = 4


ws.freeze_panes = (
    'A6'
)


# ============================================================
# 21. FORMAT DISPATCH
# ============================================================

format_data_sheet(

    wb['Dispatch'],

    'Hourly Electricity Dispatch',

    (
        'Demand, residual load, storage and '
        'generation by unit - 2023'
    ),

    NAVY
)


# ============================================================
# 22. FORMAT VPP / BOUNDARY SECTOR
# ============================================================

format_data_sheet(

    wb['VPP_Boundary'],

    'Virtual Power Plant and Boundary Sector',

    (
        'Hourly VPP electricity consumption, '
        'PowerX conversion and SectorX storage'
    ),

    GREEN
)


ws = wb[
    'VPP_Boundary'
]


# ------------------------------------------------------------
# SOC PERCENTAGE FORMAT
# ------------------------------------------------------------

soc_col = None


for cell in ws[4]:

    if (
        cell.value
        == 'SectorX Storage SOC [%]'
    ):

        soc_col = (
            cell.column
        )

        break


if soc_col:

    letter = get_column_letter(
        soc_col
    )


    for row in range(
        5,
        ws.max_row + 1
    ):

        ws.cell(
            row,
            soc_col
        ).number_format = (
            '0.0%'
        )


    ws.conditional_formatting.add(

        f'{letter}5:'
        f'{letter}{ws.max_row}',

        ColorScaleRule(

            start_type='min',
            start_color='F8696B',

            mid_type='percentile',
            mid_value=50,
            mid_color='FFEB84',

            end_type='max',
            end_color='63BE7B'
        )
    )


# ============================================================
# 23. FORMAT SYSTEM
# ============================================================

format_data_sheet(

    wb['System'],

    'System-Level Results',

    (
        'Demand, costs, emissions, prices, '
        'curtailment and reliability indicators'
    ),

    BLUE
)


ws = wb[
    'System'
]


# Highlight positive load shedding and ENS

for target_header in [

    'Load Shedding [MW]',

    'Energy Not Served [MWh]'
]:

    target_col = None


    for cell in ws[4]:

        if (
            cell.value
            == target_header
        ):

            target_col = (
                cell.column
            )

            break


    if target_col:

        letter = get_column_letter(
            target_col
        )


        ws.conditional_formatting.add(

            f'{letter}5:'
            f'{letter}{ws.max_row}',

            CellIsRule(

                operator='greaterThan',

                formula=['0'],

                fill=PatternFill(
                    'solid',
                    fgColor=LIGHT_RED
                )
            )
        )


# ============================================================
# 24. FORMAT RESERVES
# ============================================================

format_data_sheet(

    wb['Reserves'],

    'Reserve Requirements and Provision',

    (
        'Upward/downward reserve requirements, '
        'available reserves and shadow prices'
    ),

    ORANGE
)


# ============================================================
# 25. FORMAT UNITS SUMMARY
# ============================================================

format_data_sheet(

    wb['Units_Summary'],

    'Generating Units - Annual Summary',

    (
        'Installed capacity, annual generation, '
        'commitment, variable/startup costs '
        'and emissions'
    ),

    GREEN
)


ws = wb[
    'Units_Summary'
]


# ------------------------------------------------------------
# CAPACITY FACTOR PERCENTAGE
# ------------------------------------------------------------

capacity_factor_col = None


for cell in ws[4]:

    if (
        cell.value
        == 'Capacity Factor [%]'
    ):

        capacity_factor_col = (
            cell.column
        )

        break


if capacity_factor_col:

    letter = get_column_letter(
        capacity_factor_col
    )


    for row in range(
        5,
        ws.max_row + 1
    ):

        ws.cell(
            row,
            capacity_factor_col
        ).number_format = (
            '0.0%'
        )


    ws.conditional_formatting.add(

        f'{letter}5:'
        f'{letter}{ws.max_row}',

        ColorScaleRule(

            start_type='min',
            start_color='F8696B',

            mid_type='percentile',
            mid_value=50,
            mid_color='FFEB84',

            end_type='max',
            end_color='63BE7B'
        )
    )


ws.column_dimensions[
    'A'
].width = 42


# ============================================================
# 26. TAB COLORS
# ============================================================

wb[
    'Overview'
].sheet_properties.tabColor = NAVY


wb[
    'Dispatch'
].sheet_properties.tabColor = BLUE


wb[
    'VPP_Boundary'
].sheet_properties.tabColor = GREEN


wb[
    'System'
].sheet_properties.tabColor = MEDIUM_BLUE


wb[
    'Reserves'
].sheet_properties.tabColor = ORANGE


wb[
    'Units_Summary'
].sheet_properties.tabColor = GREEN


# ============================================================
# 27. SAVE - OVERWRITE SAME FINAL FILE
# ============================================================

wb.save(
    OUTPUT_FILE
)


print()
print("=" * 65)
print("FINAL EXCEL WORKBOOK UPDATED SUCCESSFULLY")
print("=" * 65)

print()
print(OUTPUT_FILE)

print()
print("Workbook sheets:")
print("  1. Overview")
print("  2. Dispatch")
print("  3. VPP_Boundary")
print("  4. System")
print("  5. Reserves")
print("  6. Units_Summary")

print()
print("Changes applied:")
print("  - Date / Time fixed in all hourly sheets")
print("  - UnitHourlyProductionCost excluded")
print("  - Average Production Cost excluded")
print("  - Revenue and Profit remain excluded")
print("  - Annual Variable Cost retained")
print("  - Annual Startup Cost retained")
print()