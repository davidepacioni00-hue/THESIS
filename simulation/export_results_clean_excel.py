import os
import pandas as pd
import dispaset as ds

from openpyxl import load_workbook
from openpyxl.styles import (
    Font, PatternFill, Border, Side, Alignment
)
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter


# ============================================================
# SETTINGS
# ============================================================

SIMULATION_PATH = r'../Simulations/simulation_(2)_BE-NUC'

OUTPUT_FILE = os.path.join(
    SIMULATION_PATH,
    'UNLEASH_BE_NUC_Clean_Results.xlsx'
)


# ============================================================
# LOAD RESULTS
# ============================================================

print()
print("=" * 60)
print("LOADING UNLEASH (2)_BE-NUC RESULTS")
print("=" * 60)

inputs, results = ds.get_sim_results(path=SIMULATION_PATH)

print("Results loaded successfully.")
print("Number of result variables:", len(results))


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def get_df(name):
    """Return a result as DataFrame and align hourly data to a common time index."""

    if name not in results:
        return None

    data = results[name]

    if isinstance(data, pd.Series):
        data = data.to_frame()

    if not isinstance(data, pd.DataFrame):
        return None

    if data.shape[1] == 0:
        return None

    df = data.copy()

    # Remove timezone because Excel cannot store timezone-aware dates
    if isinstance(df.index, pd.DatetimeIndex):
        if df.index.tz is not None:
            df.index = df.index.tz_localize(None)

    # Align all full-year hourly outputs to the same 2023 hourly index.
    # Some Dispa-SET outputs use integers instead of timestamps.
    if len(df) == 8760:
        df.index = pd.date_range(
            start='2023-01-01 00:00:00',
            periods=8760,
            freq='h'
        )

    return df

def first_column(name, new_name):
    """Return the first column of a result with a clearer name."""

    df = get_df(name)

    if df is None or df.empty:
        return None

    s = df.iloc[:, 0].copy()
    s.name = new_name

    return s


def add_series(table, series):
    """Join one Series to a DataFrame safely."""

    if series is None:
        return table

    if table is None:
        return series.to_frame()

    return table.join(series, how='outer')


def annual_sum_by_column(name):
    """Sum a result over the entire simulation period."""

    df = get_df(name)

    if df is None:
        return pd.Series(dtype=float)

    return df.sum(axis=0)


def find_unit_column(unit_df, possible_names):
    """Find a column in the units table using several possible names."""

    for name in possible_names:
        if name in unit_df.columns:
            return name

    return None


# ============================================================
# 1. UNITS
# ============================================================

if 'units' in inputs and isinstance(inputs['units'], pd.DataFrame):
    units = inputs['units'].copy()
else:
    units = pd.DataFrame()

units.index.name = 'Unit'


# ============================================================
# 2. DISPATCH SHEET
# ============================================================

dispatch = pd.DataFrame()

total_demand = first_column(
    'TotalDemand',
    'Total Demand [MW]'
)

residual_load = first_column(
    'OutputResidualLoad',
    'Residual Load [MW]'
)

dispatch = add_series(dispatch, total_demand)
dispatch = add_series(dispatch, residual_load)

power = get_df('OutputPower')

if power is not None:

    power = power.copy()

    power.columns = [
        f'Generation | {col} [MW]'
        for col in power.columns
    ]

    dispatch = dispatch.join(
        power,
        how='outer'
    )


curtailment = get_df('OutputCurtailmentPerUnit')

if curtailment is not None:

    curtailment_total = curtailment.sum(axis=1)

    curtailment_total.name = 'Total Curtailment [MW]'

    dispatch = add_series(
        dispatch,
        curtailment_total
    )


# ============================================================
# 3. VPP / BOUNDARY SECTOR SHEET
# ============================================================

vpp = pd.DataFrame()

vpp_variables = [
    ('OutputPowerX', 'VPP PowerX [MW]'),
    ('OutputPowerConsumption', 'VPP Electricity Consumption [MW]'),
    ('OutputSectorXStorageInput', 'SectorX Storage Input [MW]'),
    ('OutputSectorXStorageLevel', 'SectorX Storage Level'),
    ('OutputSectorXSelfDischarge', 'SectorX Self-Discharge'),
    ('OutputSectorXStorageShadowPrice', 'SectorX Storage Shadow Price'),
    ('SectorXShadowPrice', 'SectorX Shadow Price'),
]

for original_name, clean_name in vpp_variables:

    series = first_column(
        original_name,
        clean_name
    )

    vpp = add_series(
        vpp,
        series
    )


# ============================================================
# 4. SYSTEM SHEET
# ============================================================

system = pd.DataFrame()

system_variables = [
    ('TotalDemand', 'Total Demand [MW]'),
    ('OutputResidualLoad', 'Residual Load [MW]'),
    ('OutputSystemCost', 'System Cost'),
    ('OutputEmissions', 'Emissions'),
    ('ShadowPrice', 'Electricity Shadow Price'),
    ('OutputShedLoad', 'Load Shedding [MW]'),
    ('OutputOptimalityGap', 'Optimality Gap'),
    ('OutputOptimizationError', 'Optimization Error'),
    ('ENSH-EnergyNotServedHourly', 'Energy Not Served [MWh]'),
]

for original_name, clean_name in system_variables:

    series = first_column(
        original_name,
        clean_name
    )

    system = add_series(
        system,
        series
    )


if curtailment is not None:

    curtailment_total = curtailment.sum(axis=1)
    curtailment_total.name = 'Total Curtailment [MW]'

    system = add_series(
        system,
        curtailment_total
    )


# ============================================================
# 5. RESERVES SHEET
# ============================================================

reserves = pd.DataFrame()

reserve_variables = [
    ('OutputDemand_2U', '2U Reserve Requirement [MW]'),
    ('OutputDemand_2D', '2D Reserve Requirement [MW]'),
    ('OutputDemand_3U', '3U Reserve Requirement [MW]'),
    ('ShadowPrice_2U', '2U Reserve Shadow Price'),
    ('ShadowPrice_2D', '2D Reserve Shadow Price'),
    ('ShadowPrice_3U', '3U Reserve Shadow Price'),
]

for original_name, clean_name in reserve_variables:

    series = first_column(
        original_name,
        clean_name
    )

    reserves = add_series(
        reserves,
        series
    )


for source_name, clean_name in [
    ('OutputReserve_2U', 'Available 2U Reserve [MW]'),
    ('OutputReserve_2D', 'Available 2D Reserve [MW]'),
    ('OutputReserve_3U', 'Available 3U Reserve [MW]'),
]:

    reserve_df = get_df(source_name)

    if reserve_df is not None:

        total = reserve_df.sum(axis=1)
        total.name = clean_name

        reserves = add_series(
            reserves,
            total
        )


# ============================================================
# 6. UNIT ECONOMICS
# ============================================================

generation = annual_sum_by_column(
    'OutputPower'
)

revenue = annual_sum_by_column(
    'UnitHourlyRevenue'
)

power_revenue = annual_sum_by_column(
    'UnitHourlyPowerRevenue'
)

production_cost = annual_sum_by_column(
    'UnitHourlyProductionCost'
)

startup_cost = annual_sum_by_column(
    'UnitHourlyStartUpCost'
)

variable_cost = annual_sum_by_column(
    'UnitHourlyVariableCost'
)

profit = annual_sum_by_column(
    'UnitHourlyProfit'
)

committed = annual_sum_by_column(
    'OutputCommitted'
)


all_units = sorted(
    set(generation.index)
    | set(revenue.index)
    | set(production_cost.index)
    | set(profit.index)
)


unit_economics = pd.DataFrame(
    index=all_units
)

unit_economics.index.name = 'Unit'

unit_economics['Annual Generation [MWh]'] = generation
unit_economics['Committed Hours'] = committed
unit_economics['Total Revenue'] = revenue
unit_economics['Power Revenue'] = power_revenue
unit_economics['Variable Cost'] = variable_cost
unit_economics['Startup Cost'] = startup_cost
unit_economics['Production Cost'] = production_cost
unit_economics['Profit'] = profit

unit_economics = unit_economics.fillna(0)


# Add some unit characteristics when available

if not units.empty:

    fuel_col = find_unit_column(
        units,
        ['Fuel', 'fuel']
    )

    technology_col = find_unit_column(
        units,
        ['Technology', 'technology']
    )

    zone_col = find_unit_column(
        units,
        ['Zone', 'zone']
    )

    capacity_col = find_unit_column(
        units,
        [
            'PowerCapacity',
            'Power Capacity',
            'Capacity'
        ]
    )

    if zone_col:
        unit_economics.insert(
            0,
            'Zone',
            units[zone_col].reindex(unit_economics.index)
        )

    if technology_col:
        unit_economics.insert(
            1 if zone_col else 0,
            'Technology',
            units[technology_col].reindex(unit_economics.index)
        )

    if fuel_col:
        unit_economics.insert(
            2 if zone_col and technology_col else 0,
            'Fuel',
            units[fuel_col].reindex(unit_economics.index)
        )

    if capacity_col:

        position = min(
            3,
            len(unit_economics.columns)
        )

        unit_economics.insert(
            position,
            'Installed Capacity [MW]',
            units[capacity_col].reindex(unit_economics.index)
        )


# ============================================================
# 7. OVERVIEW / KPIs
# ============================================================

hours = 0

if total_demand is not None:
    hours = len(total_demand)


def safe_sum(series):
    if series is None:
        return 0
    return float(series.fillna(0).sum())


annual_demand = safe_sum(
    total_demand
)

system_cost = safe_sum(
    first_column(
        'OutputSystemCost',
        'System Cost'
    )
)

emissions = safe_sum(
    first_column(
        'OutputEmissions',
        'Emissions'
    )
)

shed_load = safe_sum(
    first_column(
        'OutputShedLoad',
        'Load Shedding'
    )
)

energy_not_served = safe_sum(
    first_column(
        'ENSH-EnergyNotServedHourly',
        'Energy Not Served'
    )
)


if curtailment is not None:
    annual_curtailment = float(
        curtailment.fillna(0).sum().sum()
    )
else:
    annual_curtailment = 0


def scalar_result(name):

    df = get_df(name)

    if df is None or df.empty:
        return 0

    try:
        return float(
            df.iloc[0, 0]
        )
    except Exception:
        return 0


eens = scalar_result(
    'EENS-ExpectedEnergyNotServed'
)

lolh = scalar_result(
    'LOLH-LosOfLoadHours'
)

lolp = scalar_result(
    'LOLP-LosOfLoadProbability'
)

lolf = scalar_result(
    'LOLF-LosOfLoadFrequency'
)


overview_data = [
    ['Scenario', '(2)_BE-NUC'],
    ['Model', 'Dispa-SET UNLEASH Boundary Sector'],
    ['Country / Zone', 'Belgium (BE)'],
    ['Simulation Year', '2023'],
    ['Simulated Hours', hours],
    ['Annual Electricity Demand [MWh]', annual_demand],
    ['Total System Cost', system_cost],
    ['Total Emissions', emissions],
    ['Total Curtailment [MWh]', annual_curtailment],
    ['Total Load Shedding [MWh]', shed_load],
    ['Energy Not Served [MWh]', energy_not_served],
    ['Expected Energy Not Served', eens],
    ['Loss of Load Hours', lolh],
    ['Loss of Load Probability', lolp],
    ['Loss of Load Frequency', lolf],
]

overview = pd.DataFrame(
    overview_data,
    columns=['Indicator', 'Value']
)


# ============================================================
# EXPORT TO EXCEL
# ============================================================

print()
print("Creating clean Excel workbook...")

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine='openpyxl'
) as writer:

    overview.to_excel(
        writer,
        sheet_name='Overview',
        index=False,
        startrow=3
    )

    units.to_excel(
        writer,
        sheet_name='Units'
    )

    dispatch.to_excel(
        writer,
        sheet_name='Dispatch'
    )

    vpp.to_excel(
        writer,
        sheet_name='VPP_Boundary'
    )

    system.to_excel(
        writer,
        sheet_name='System'
    )

    reserves.to_excel(
        writer,
        sheet_name='Reserves'
    )

    unit_economics.to_excel(
        writer,
        sheet_name='Unit_Economics'
    )


# ============================================================
# PROFESSIONAL FORMATTING
# ============================================================

print("Applying professional formatting...")

wb = load_workbook(
    OUTPUT_FILE
)

# Color palette

DARK_BLUE = '17365D'
BLUE = '1F4E78'
MEDIUM_BLUE = '5B9BD5'
LIGHT_BLUE = 'D9EAF7'
VERY_LIGHT_BLUE = 'EEF5FB'

DARK_GREEN = '548235'
LIGHT_GREEN = 'E2F0D9'

ORANGE = 'ED7D31'
LIGHT_ORANGE = 'FCE4D6'

GRAY = 'D9E1F2'
LIGHT_GRAY = 'F2F2F2'

WHITE = 'FFFFFF'
BLACK = '000000'

thin_gray = Side(
    style='thin',
    color='D9E1F2'
)


def style_standard_sheet(
    ws,
    freeze='B2',
    auto_filter=True
):

    ws.freeze_panes = freeze

    max_row = ws.max_row
    max_col = ws.max_column

    # Header

    for cell in ws[1]:

        cell.fill = PatternFill(
            'solid',
            fgColor=DARK_BLUE
        )

        cell.font = Font(
            color=WHITE,
            bold=True
        )

        cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

        cell.border = Border(
            bottom=thin_gray
        )

    ws.row_dimensions[1].height = 28

    # Alternating rows

    for row in range(
        2,
        max_row + 1
    ):

        if row % 2 == 0:

            fill = PatternFill(
                'solid',
                fgColor=VERY_LIGHT_BLUE
            )

            for col in range(
                1,
                max_col + 1
            ):
                ws.cell(
                    row,
                    col
                ).fill = fill

    # Borders and alignment

    for row in ws.iter_rows():

        for cell in row:

            cell.border = Border(
                bottom=thin_gray
            )

            cell.alignment = Alignment(
                vertical='center'
            )

    # Auto filter

    if auto_filter:
        ws.auto_filter.ref = ws.dimensions

    # Column widths

    for col_idx in range(
        1,
        max_col + 1
    ):

        letter = get_column_letter(
            col_idx
        )

        max_length = 0

        for cell in ws[letter]:

            try:
                text = str(
                    cell.value
                    if cell.value is not None
                    else ''
                )

                max_length = max(
                    max_length,
                    len(text)
                )

            except Exception:
                pass

        width = min(
            max(max_length + 2, 12),
            38
        )

        ws.column_dimensions[
            letter
        ].width = width

    # Numeric formats

    for row in ws.iter_rows(
        min_row=2
    ):

        for cell in row:

            if isinstance(
                cell.value,
                (int, float)
            ):

                cell.number_format = (
                    '#,##0.00'
                )


# ------------------------------------------------------------
# OVERVIEW
# ------------------------------------------------------------

ws = wb['Overview']

ws.merge_cells(
    'A1:D1'
)

title = ws['A1']

title.value = (
    'UNLEASH – (2)_BE-NUC Simulation Results'
)

title.fill = PatternFill(
    'solid',
    fgColor=DARK_BLUE
)

title.font = Font(
    color=WHITE,
    bold=True,
    size=18
)

title.alignment = Alignment(
    horizontal='center',
    vertical='center'
)

ws.row_dimensions[1].height = 36


ws.merge_cells(
    'A2:D2'
)

subtitle = ws['A2']

subtitle.value = (
    'Belgian Boundary Sector / Virtual Power Plant – 2023'
)

subtitle.fill = PatternFill(
    'solid',
    fgColor=MEDIUM_BLUE
)

subtitle.font = Font(
    color=WHITE,
    italic=True,
    size=11
)

subtitle.alignment = Alignment(
    horizontal='center'
)


# Overview headers start at row 4

for cell in ws[4]:

    if cell.column <= 2:

        cell.fill = PatternFill(
            'solid',
            fgColor=BLUE
        )

        cell.font = Font(
            color=WHITE,
            bold=True
        )


for row in range(
    5,
    ws.max_row + 1
):

    indicator_cell = ws.cell(
        row,
        1
    )

    value_cell = ws.cell(
        row,
        2
    )

    indicator_cell.fill = PatternFill(
        'solid',
        fgColor=LIGHT_BLUE
    )

    indicator_cell.font = Font(
        bold=True,
        color=DARK_BLUE
    )

    value_cell.fill = PatternFill(
        'solid',
        fgColor=WHITE
    )

    if isinstance(
        value_cell.value,
        (int, float)
    ):
        value_cell.number_format = '#,##0.00'

    indicator_cell.border = Border(
        bottom=thin_gray
    )

    value_cell.border = Border(
        bottom=thin_gray
    )


ws.column_dimensions['A'].width = 38
ws.column_dimensions['B'].width = 28

ws.freeze_panes = 'A5'


# ------------------------------------------------------------
# STANDARD SHEETS
# ------------------------------------------------------------

for sheet_name in [
    'Units',
    'Dispatch',
    'VPP_Boundary',
    'System',
    'Reserves',
    'Unit_Economics'
]:

    style_standard_sheet(
        wb[sheet_name]
    )


# ------------------------------------------------------------
# SPECIAL COLORS: VPP
# ------------------------------------------------------------

ws = wb['VPP_Boundary']

for cell in ws[1]:

    cell.fill = PatternFill(
        'solid',
        fgColor=DARK_GREEN
    )


# ------------------------------------------------------------
# SPECIAL COLORS: SYSTEM
# ------------------------------------------------------------

ws = wb['System']

for cell in ws[1]:

    cell.fill = PatternFill(
        'solid',
        fgColor=BLUE
    )


# ------------------------------------------------------------
# SPECIAL COLORS: RESERVES
# ------------------------------------------------------------

ws = wb['Reserves']

for cell in ws[1]:

    cell.fill = PatternFill(
        'solid',
        fgColor=ORANGE
    )


# ------------------------------------------------------------
# CONDITIONAL FORMATTING: UNIT PROFIT
# ------------------------------------------------------------

ws = wb['Unit_Economics']

profit_col = None

for cell in ws[1]:

    if cell.value == 'Profit':

        profit_col = cell.column

        break


if profit_col is not None:

    letter = get_column_letter(
        profit_col
    )

    ws.conditional_formatting.add(
        f'{letter}2:{letter}{ws.max_row}',
        ColorScaleRule(
            start_type='min',
            start_color='F8696B',
            mid_type='percentile',
            mid_value=50,
            mid_color='FFEB84',
            end_type='max',
            end_color='63BE7B'
        )
    )


# ------------------------------------------------------------
# TAB COLORS
# ------------------------------------------------------------

wb['Overview'].sheet_properties.tabColor = DARK_BLUE
wb['Units'].sheet_properties.tabColor = MEDIUM_BLUE
wb['Dispatch'].sheet_properties.tabColor = BLUE
wb['VPP_Boundary'].sheet_properties.tabColor = DARK_GREEN
wb['System'].sheet_properties.tabColor = MEDIUM_BLUE
wb['Reserves'].sheet_properties.tabColor = ORANGE
wb['Unit_Economics'].sheet_properties.tabColor = DARK_GREEN


# ------------------------------------------------------------
# SAVE
# ------------------------------------------------------------

wb.save(
    OUTPUT_FILE
)


print()
print("=" * 60)
print("CLEAN EXCEL WORKBOOK CREATED SUCCESSFULLY")
print("=" * 60)

print()
print(OUTPUT_FILE)

print()
print("Workbook sheets:")
print("  1. Overview")
print("  2. Units")
print("  3. Dispatch")
print("  4. VPP_Boundary")
print("  5. System")
print("  6. Reserves")
print("  7. Unit_Economics")
print()